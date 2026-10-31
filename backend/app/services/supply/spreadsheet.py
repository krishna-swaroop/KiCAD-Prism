"""The first sheet of an Excel file as rows of text. Pure apart from openpyxl; no database.

Used for uploaded BOMs and the ECHA candidate list. Only Office Open XML files
(.xlsx, .xlsm) are read; the archive's unpacked size is checked before openpyxl
opens it, so a small file cannot unpack into gigabytes.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from app.services.supply.errors import Refused


class SpreadsheetError(Refused):
    """The file is too large or not a readable spreadsheet. The message is for the person who chose it."""

    code = "invalid_spreadsheet"


def _text(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.date().isoformat()
    if isinstance(cell, date):
        return cell.isoformat()
    if isinstance(cell, float) and cell.is_integer():
        # Excel stores every number as a float; a whole one shows without ".0".
        return str(int(cell))
    return str(cell)


def first_sheet_rows(data: bytes, *, max_bytes: int, max_unpacked: int) -> list[list[str]]:
    """Every non-blank row of the first sheet, each cell as text."""

    from openpyxl import load_workbook

    if len(data) > max_bytes:
        raise SpreadsheetError(f"The file is larger than {max_bytes // (1024 * 1024)} MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > max_unpacked:
                raise SpreadsheetError("The spreadsheet is too large once unpacked.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            rows = [[_text(cell) for cell in row] for row in book.worksheets[0].iter_rows(values_only=True)]
        finally:
            book.close()
    except SpreadsheetError:
        raise
    except Exception as exc:  # noqa: BLE001 - a damaged or foreign file is the person's to fix, whatever the library raised
        raise SpreadsheetError("The file is not a readable Excel spreadsheet. Export it as CSV, or save it again as .xlsx.") from exc
    return [row for row in rows if any(cell.strip() for cell in row)]
